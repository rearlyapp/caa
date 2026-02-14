#!/usr/bin/env python3
"""Local CPU backend for document extraction and Excel generation."""

from __future__ import annotations

import base64
import io
import json
import os
import re
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Literal

import openpyxl
import torch
from fastapi import FastAPI, HTTPException
from PIL import Image, ImageEnhance, ImageOps
from pydantic import BaseModel, Field
from transformers import AutoProcessor, Qwen3VLForConditionalGeneration

try:
    import pypdfium2 as pdfium
except ImportError:  # pragma: no cover - optional at import, required for PDF input
    pdfium = None

MODEL_DIR = Path(os.getenv("LOCAL_MODEL_DIR", "./models/Qwen3-VL-2B-Instruct"))
DEFAULT_TEMPLATE_PATH = os.getenv(
    "CHECKLIST_TEMPLATE_PATH", "Checklist for Company Formation dated 11 02 2026.xlsx"
)
MAX_PIXELS = 4_000_000
MAX_SIDE = 1800

PAN_PROMPT = """Look at this Indian PAN card image carefully.

On a PAN card, the layout is:
- Line 1: Cardholder's FULL NAME (e.g. "KISHORE SHEIK AHAMED M R")
- Line 2: Father's FULL NAME (e.g. "MOHAMMED RAHAMATHULLA")
- Line 3: Date of Birth (DD/MM/YYYY)
- Bottom: 10-character PAN number (e.g. BQJPK6347Q)

Extract these fields as JSON. The name and father's name are SEPARATE people on SEPARATE lines:

{
  "name": "cardholder name only (first name line)",
  "fathers_name": "father's name only (second name line)",
  "date_of_birth": "DD/MM/YYYY",
  "pan_number": "10 character PAN"
}

Return ONLY the JSON object, no other text."""

AADHAAR_PROMPT = """Look at this Aadhaar card image carefully.
Extract ALL the following fields and return ONLY valid JSON, nothing else:

{
  "name": "full name as printed",
  "aadhaar_number": "12 digit number (XXXX XXXX XXXX format)",
  "date_of_birth": "DD/MM/YYYY format",
  "gender": "MALE or FEMALE",
  "address": "complete residential address as printed including pincode"
}

Rules:
- Copy text EXACTLY as printed
- Aadhaar number is exactly 12 digits
- Include full address with pincode
- Return ONLY the JSON object, no other text"""

PAN_PATTERN = re.compile(r"[A-Z]{5}[0-9]{4}[A-Z]")
AADHAAR_PATTERN = re.compile(r"([0-9]{4})\s?([0-9]{4})\s?([0-9]{4})")


class ExtractRequest(BaseModel):
    image_base64: str = Field(min_length=32)
    document_type: Literal["pan", "aadhaar"]
    file_name: str | None = None
    mime_type: str | None = None


class ExportRequest(BaseModel):
    case_name: str = Field(min_length=1)
    case_data: dict[str, Any]
    template_path: str | None = None


class ModelRuntime:
    def __init__(self) -> None:
        self.model: Qwen3VLForConditionalGeneration | None = None
        self.processor: AutoProcessor | None = None
        self.load_error: str | None = None
        self.load_seconds: float | None = None

    def load(self) -> None:
        if self.model is not None and self.processor is not None:
            return

        if not MODEL_DIR.exists():
            raise FileNotFoundError(
                f"Local model directory not found: {MODEL_DIR}. "
                "Set LOCAL_MODEL_DIR to a valid Qwen3-VL-2B-Instruct path."
            )

        start = time.time()
        self.model = Qwen3VLForConditionalGeneration.from_pretrained(
            str(MODEL_DIR),
            dtype=torch.float32,
        )
        self.processor = AutoProcessor.from_pretrained(str(MODEL_DIR))
        self.load_seconds = time.time() - start


runtime = ModelRuntime()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    try:
        runtime.load()
    except Exception as exc:
        runtime.load_error = str(exc)
    yield


app = FastAPI(title="CAA Local Backend", lifespan=lifespan)


def parse_json_output(text: str) -> dict[str, Any]:
    text = text.strip()
    if not text:
        return {}

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(1))
        except json.JSONDecodeError:
            pass

    match = re.search(r"\{.*\}", text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(0))
        except json.JSONDecodeError:
            pass

    return {}


def render_pdf_first_page(pdf_bytes: bytes) -> Image.Image:
    if pdfium is None:
        raise ValueError("PDF conversion requires pypdfium2 to be installed")

    document = pdfium.PdfDocument(pdf_bytes)
    if len(document) < 1:
        raise ValueError("PDF has no pages")

    page = document[0]
    bitmap = page.render(scale=2.0)
    image = bitmap.to_pil().convert("RGB")
    return image


def decode_base64_image(
    image_base64: str, file_name: str | None = None, mime_type: str | None = None
) -> Image.Image:
    inferred_mime = mime_type
    payload = image_base64

    if image_base64.startswith("data:") and "," in image_base64:
        header, payload = image_base64.split(",", 1)
        if ";" in header:
            inferred_mime = header[5:].split(";", 1)[0]

    image_bytes = base64.b64decode(payload)
    lower_name = (file_name or "").lower()
    is_pdf = inferred_mime == "application/pdf" or lower_name.endswith(".pdf")

    if is_pdf:
        return render_pdf_first_page(image_bytes)

    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    return image


def preprocess_image(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image).convert("RGB")
    width, height = image.size

    side_scale = min(1.0, MAX_SIDE / max(width, height))
    pixel_scale = min(1.0, (MAX_PIXELS / (width * height)) ** 0.5)
    scale = min(side_scale, pixel_scale)

    if scale < 1.0:
        resized = (max(1, int(width * scale)), max(1, int(height * scale)))
        image = image.resize(resized, Image.Resampling.LANCZOS)

    image = ImageOps.autocontrast(image, cutoff=1)
    image = ImageEnhance.Contrast(image).enhance(1.15)
    return image


def normalize_pan(data: dict[str, Any]) -> dict[str, str]:
    pan_number = str(data.get("pan_number", "")).upper().strip().replace(" ", "")
    match = PAN_PATTERN.search(pan_number)
    pan_number = match.group(0) if match else pan_number

    return {
        "name": str(data.get("name", "")).strip(),
        "fathers_name": str(data.get("fathers_name", "")).strip(),
        "date_of_birth": str(data.get("date_of_birth", "")).strip(),
        "pan_number": pan_number,
    }


def normalize_aadhaar(data: dict[str, Any]) -> dict[str, str]:
    aadhaar = str(data.get("aadhaar_number", "")).strip()
    digits = re.sub(r"\D", "", aadhaar)
    if len(digits) == 12:
        aadhaar = f"{digits[:4]} {digits[4:8]} {digits[8:]}"
    else:
        match = AADHAAR_PATTERN.search(aadhaar)
        if match:
            aadhaar = " ".join(match.groups())

    return {
        "name": str(data.get("name", "")).strip(),
        "aadhaar_number": aadhaar,
        "date_of_birth": str(data.get("date_of_birth", "")).strip(),
        "gender": str(data.get("gender", "")).strip().upper(),
        "address": str(data.get("address", "")).strip(),
    }


def extract_with_qwen(image: Image.Image, document_type: Literal["pan", "aadhaar"]) -> dict[str, Any]:
    if runtime.load_error:
        raise HTTPException(
            status_code=503, detail=f"Model is unavailable: {runtime.load_error}"
        )
    if runtime.model is None or runtime.processor is None:
        raise HTTPException(status_code=503, detail="Model is still loading")

    prompt = PAN_PROMPT if document_type == "pan" else AADHAAR_PROMPT
    messages = [
        {
            "role": "user",
            "content": [
                {"type": "image", "image": image},
                {"type": "text", "text": prompt},
            ],
        }
    ]

    inputs = runtime.processor.apply_chat_template(
        messages,
        tokenize=True,
        add_generation_prompt=True,
        return_dict=True,
        return_tensors="pt",
    )
    inputs = inputs.to(runtime.model.device)

    start = time.time()
    with torch.no_grad():
        generated_ids = runtime.model.generate(**inputs, max_new_tokens=512)
    elapsed = time.time() - start

    generated_ids_trimmed = [
        out_ids[len(in_ids) :] for in_ids, out_ids in zip(inputs.input_ids, generated_ids)
    ]
    raw_text = runtime.processor.batch_decode(
        generated_ids_trimmed,
        skip_special_tokens=True,
        clean_up_tokenization_spaces=False,
    )[0]

    parsed = parse_json_output(raw_text)
    normalized = normalize_pan(parsed) if document_type == "pan" else normalize_aadhaar(parsed)

    return {
        "document_type": document_type,
        "raw_text": raw_text,
        "elapsed": round(elapsed, 3),
        "extracted_data": normalized,
    }


def set_cell_if_value(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int, value: Any) -> None:
    if value is None:
        return
    value_str = str(value).strip()
    if value_str:
        ws.cell(row=row, column=col, value=value_str)


def document_uploaded(documents: list[dict[str, Any]] | None, doc_type: str) -> str:
    if not documents:
        return ""
    for doc in documents:
        if doc.get("type") == doc_type and doc.get("fileName"):
            return "Uploaded"
    return ""


def populate_director(ws: openpyxl.worksheet.worksheet.Worksheet, director: dict[str, Any], col: int) -> None:
    pan = director.get("panData") or {}
    aadhaar = director.get("aadhaarData") or {}
    docs = director.get("documents") or []

    set_cell_if_value(ws, 4, col, pan.get("name"))
    set_cell_if_value(ws, 5, col, pan.get("fathers_name"))
    set_cell_if_value(ws, 6, col, pan.get("date_of_birth"))
    set_cell_if_value(ws, 7, col, director.get("placeOfBirth"))
    set_cell_if_value(ws, 8, col, director.get("nationality") or "Indian")
    set_cell_if_value(ws, 9, col, director.get("residentOfIndia") or "Yes")
    set_cell_if_value(ws, 10, col, director.get("occupation"))
    set_cell_if_value(ws, 11, col, director.get("education"))
    set_cell_if_value(ws, 12, col, director.get("sharesSubscribed"))
    set_cell_if_value(ws, 13, col, director.get("durationAtAddress"))
    set_cell_if_value(ws, 14, col, director.get("email"))
    set_cell_if_value(ws, 15, col, director.get("mobile"))
    set_cell_if_value(ws, 16, col, pan.get("pan_number"))
    aadhaar_number = aadhaar.get("aadhaar_number")
    if aadhaar_number:
        set_cell_if_value(ws, 17, col, f"AADHAR - {aadhaar_number}")
    set_cell_if_value(ws, 18, col, aadhaar.get("address"))
    set_cell_if_value(ws, 19, col, director.get("dinNumber"))
    set_cell_if_value(ws, 20, col, document_uploaded(docs, "photo"))
    set_cell_if_value(ws, 21, col, document_uploaded(docs, "signature"))


def populate_company(ws: openpyxl.worksheet.worksheet.Worksheet, company: dict[str, Any]) -> None:
    set_cell_if_value(ws, 24, 3, "Uploaded" if company.get("electricityBillUploaded") else "")
    set_cell_if_value(ws, 25, 3, company.get("latitude"))
    set_cell_if_value(ws, 26, 3, company.get("longitude"))
    set_cell_if_value(ws, 27, 3, "Uploaded" if company.get("nocUploaded") else "")
    set_cell_if_value(ws, 28, 3, company.get("officeEmail"))
    set_cell_if_value(ws, 29, 3, company.get("officeMobile"))
    set_cell_if_value(ws, 30, 3, company.get("authorizedShareCapital"))
    set_cell_if_value(ws, 31, 3, company.get("paidUpShareCapital"))
    set_cell_if_value(ws, 32, 3, company.get("objectives"))
    set_cell_if_value(ws, 33, 3, company.get("otherObjectives"))


def populate_professional(ws: openpyxl.worksheet.worksheet.Worksheet, professional: dict[str, Any]) -> None:
    set_cell_if_value(ws, 35, 3, professional.get("name"))
    set_cell_if_value(ws, 36, 3, professional.get("membershipNo"))
    set_cell_if_value(ws, 37, 3, professional.get("address"))


def resolve_template_path(request_path: str | None) -> Path:
    if request_path:
        candidate = Path(request_path)
        if candidate.exists():
            return candidate

    configured = Path(DEFAULT_TEMPLATE_PATH)
    if configured.exists():
        return configured

    fallback = sorted(Path(".").glob("Checklist for Company Formation dated *.xlsx"))
    if fallback:
        return fallback[-1]

    raise FileNotFoundError(
        "Checklist template not found. Provide template_path or set CHECKLIST_TEMPLATE_PATH."
    )


def safe_file_name(case_name: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "_", case_name).strip("_")
    if not normalized:
        normalized = "case"
    return f"{normalized}_Checklist_filled.xlsx"


@app.get("/health")
def health() -> dict[str, Any]:
    if runtime.load_error:
        return {
            "status": "degraded",
            "model_loaded": False,
            "error": runtime.load_error,
            "model_dir": str(MODEL_DIR),
        }

    return {
        "status": "ok",
        "model_loaded": runtime.model is not None,
        "model_dir": str(MODEL_DIR),
        "model_load_seconds": runtime.load_seconds,
    }


@app.post("/extract")
def extract(request: ExtractRequest) -> dict[str, Any]:
    try:
        image = decode_base64_image(
            request.image_base64, file_name=request.file_name, mime_type=request.mime_type
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid document payload: {exc}") from exc

    image = preprocess_image(image)
    return extract_with_qwen(image, request.document_type)


@app.post("/export")
def export(request: ExportRequest) -> dict[str, Any]:
    try:
        template_path = resolve_template_path(request.template_path)
        workbook = openpyxl.load_workbook(template_path)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Unable to load template: {exc}") from exc

    ws = workbook.active
    directors = request.case_data.get("directors") or []

    if directors:
        populate_director(ws, directors[0] if len(directors) > 0 else {}, 3)
        populate_director(ws, directors[1] if len(directors) > 1 else {}, 4)

    populate_company(ws, request.case_data.get("companyInfo") or {})
    populate_professional(ws, request.case_data.get("professionalInfo") or {})

    output = io.BytesIO()
    workbook.save(output)
    encoded = base64.b64encode(output.getvalue()).decode("utf-8")

    return {
        "success": True,
        "file_name": safe_file_name(request.case_name),
        "file_content_base64": encoded,
        "template_used": str(template_path),
    }
